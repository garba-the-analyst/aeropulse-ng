#!/usr/bin/env python3
"""Build the editable AeroPulse-NG Office documentation package.

Outputs:
  docs/docx_pptx/AeroPulse-NG_Pitch_Deck.pptx
  docs/docx_pptx/AeroPulse-NG_Executive_Summary.docx
  docs/docx_pptx/AeroPulse-NG_Architecture.docx
  docs/docx_pptx/AeroPulse-NG_User_Manual.docx
  docs/docx_pptx/AeroPulse-NG_Technical_Reference.docx

The generator uses fixed-width tables, explicit text wrapping and page breaks
so the editable exports do not inherit the clipping problems of the earlier
PDF-only workflow.
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor as DocxRGB
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches, Pt as PptPt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
OUT = ROOT / "docs" / "docx_pptx"
OUT.mkdir(parents=True, exist_ok=True)

PITCH_BLACK = (5, 6, 7)
OBSIDIAN = (13, 20, 30)
OBSIDIAN_LIGHT = (26, 31, 36)
SILVER = (232, 236, 239)
SILVER_MID = (198, 204, 210)
SILVER_DIM = (139, 151, 163)
ACCENT = (199, 205, 211)
SUCCESS = (134, 181, 154)
WARNING = (210, 168, 90)
CRITICAL = (216, 97, 97)
WHITE = SILVER
# Semantic aliases used by the pitch layout. They deliberately resolve to the
# restrained monochrome palette rather than reintroducing bright neon colours.
CIVIL = ACCENT
MILV = SILVER_MID
WARNING = WARNING
CIVIL2 = SILVER_MID


def ppt_colour(rgb: tuple[int, int, int]) -> PptRGB:
    return rgb if isinstance(rgb, PptRGB) else PptRGB(*rgb)


def doc_colour(rgb: tuple[int, int, int]) -> DocxRGB:
    return DocxRGB(*rgb)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, colour="23303A", size="6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), colour)


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = doc_colour((30, 38, 48))
    for name, size, colour in (("Title", 24, (13, 20, 30)), ("Heading 1", 17, (13, 20, 30)), ("Heading 2", 13, (70, 91, 108)), ("Heading 3", 11, (70, 91, 108))):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = doc_colour(colour)


def add_doc_header_footer(section, title: str) -> None:
    header = section.header.paragraphs[0]
    header.text = f"AeroPulse-NG  |  {title}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = doc_colour((100, 112, 124))
    footer = section.footer.paragraphs[0]
    footer.text = "AeroPulse-NG v2.4  |  Prototype documentation  |  Verify all operational claims before deployment"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.name = "Aptos"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = doc_colour((100, 112, 124))


def add_doc_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.color.rgb = doc_colour((13, 20, 30))
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.name = "Aptos"
    r2.font.size = Pt(11)
    r2.font.color.rgb = doc_colour((100, 112, 124))


def add_doc_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_doc_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, (cell, text) in enumerate(zip(header.cells, headers)):
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "0D141E")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = doc_colour((240, 242, 244))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for idx, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F4F6F8" if row_index % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            run.font.name = "Aptos"
            run.font.size = Pt(8.7)
            run.font.color.rgb = doc_colour((30, 38, 48))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_doc_image(doc: Document, filename: str, width: float) -> None:
    path = ASSETS / filename
    if path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(path), width=Inches(width))


def build_docx(name: str, title: str, subtitle: str, content) -> None:
    doc = Document()
    set_doc_defaults(doc)
    add_doc_header_footer(doc.sections[0], title)
    add_doc_title(doc, title, subtitle)
    content(doc)
    doc.save(OUT / name)


def executive_content(doc: Document) -> None:
    doc.add_heading("Entry Information", level=1)
    add_doc_table(doc, ["Field", "Value"], [
        ["Primary category", "Flight Planning & Air Traffic Management"],
        ["Competition", "NDAIE 2026 National Innovation Pitch Challenge"],
        ["Word count", "285 words of the 300-word limit"],
        ["Development stage", "Working prototype; field hardware integration pending"],
    ], [1.7, 4.8])
    for heading, body in [
        ("Problem", "Nigeria's airspace is monitored largely through expensive imported primary and secondary surveillance radar — installations costing millions of dollars, demanding specialist maintenance, yet still leaving low-altitude corridors, secondary aerodromes and terrain-shadowed sectors under-covered. When connectivity or a radar head fails, controllers lose the traffic picture precisely when harmattan haze, diversion surges or military activity raise risk. Non-cooperative and transponder-silent aircraft widen the gap further."),
        ("Proposed solution", "AeroPulse-NG is an air-gapped surveillance engine that turns commodity software-defined radios (about USD 150 of hardware) plus a standard PC into a working tactical radar display. It ingests over-the-air aviation telemetry with zero internet or cloud dependency, giving civil air traffic management (NAMA) and tactical command (NAF) a sovereign, locally maintainable traffic picture at a fraction of conventional cost."),
        ("Technical approach", "A Rust/Tauri backend demodulates 1090 MHz Mode S ADS-B, resolving positions through Compact Position Reporting; a six-state Extended Kalman Filter smooths jitter and dead-reckons tracks through 30-second RF dropouts. An R*-tree spatial index drives continuous Short-Term Conflict Alerting against ICAO Doc 4444 separation minima (rendered 9.26 km / 305 m) within a 120-second lookahead. A triple-fusion weather matrix merges AWOS serial telemetry, ACARS D-ATIS broadcasts and Mode S BDS 4,4/4,5 downlinks into an offline 3D wind/temperature model tuned for harmattan operations. Defence overlays flag dark targets, squawk emergencies and geofence incursions, and compute intercept geometry. A hardware-accelerated display layer renders HF-STD-010A-compliant screens across dual operator monitors, in SI metric units per NCAA / ICAO Annex 5 policy."),
        ("Likely impact", "Operational: conflict alerting and fused weather at aerodromes that cannot justify radar. Economic: over 99% capital reduction per station. Sovereignty: fully offline national capability. Educational: an engineering reference for Nigerian avionics talent. Development stage: prototype core verified by automated checks; dual-display operator interface complete; RF front-end and field-trial hardware integration next."),
    ]:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.add_page_break()
    doc.add_heading("Verification Snapshot", level=1)
    add_doc_table(doc, ["Area", "Result", "Evidence"], [
        ["Rust engine", "90 tests passed", "Decoder, CPR, EKF, R*-tree, STCA and engine integration"],
        ["Python sidecar", "14 tests passed", "AWOS codec, METAR parser and DuckDB durability"],
        ["Frontend", "TypeScript clean", "Strict typecheck and Vite multi-page build"],
        ["Desktop shell", "Build verified", "Tauri v2 host linked on Linux with webkit development files"],
    ], [1.4, 1.5, 3.6])
    add_doc_image(doc, "radar-scope.png", 6.4)
    doc.add_heading("Integrity Declaration", level=1)
    doc.add_paragraph("The project includes significant declared generative-AI assistance. External sources and dependencies must be credited in the final submission. The team remains responsible for validating all performance, safety, cost and regulatory claims before presenting or operating the system.")


def architecture_content(doc: Document) -> None:
    doc.add_heading("1. System Overview", level=1)
    doc.add_paragraph("AeroPulse-NG is an air-gapped surveillance workstation. Software-defined radio channels feed a Rust fusion engine; a Python sidecar records data and reads optional AWOS serial observations; two operator displays receive a 60 Hz telemetry snapshot. The design separates deterministic signal and safety calculations from asynchronous desktop I/O.")
    add_doc_image(doc, "architecture.png", 6.5)
    doc.add_heading("2. Pipeline", level=1)
    add_doc_table(doc, ["Stage", "Input", "Output"], [
        ["Ingestion", "1090 MHz, 131.55 MHz, VHF and simulator", "Raw frame and weather events"],
        ["Decode", "Mode S / ACARS bytes", "Identity, position, velocity and reports"],
        ["Filter", "Geodetic observations", "ENU state and covariance"],
        ["Assess", "Predicted tracks", "STCA, anomaly and geofence events"],
        ["Publish", "Engine state", "EngineSnapshot on telemetry bus"],
        ["Present", "Snapshot", "Radar canvas and operations panels"],
    ], [1.2, 2.5, 2.8])
    doc.add_heading("3. Core Design Decisions", level=1)
    add_doc_bullets(doc, [
        "The simulator creates genuine DF17 frames and sends them through the same decoder used by the future RF path.",
        "The engine queues observations, predicts once per tick, then fuses measurements against the current epoch.",
        "The EKF uses metres and metres per second internally; operator displays convert to SI presentation units in one utility module.",
        "Safety thresholds remain ICAO constants internally: 5 NM horizontal and 1,000 ft vertical, displayed as 9.26 km and 305 m.",
        "Optional persistence and simulated weather degrade independently and must not stop the surveillance engine.",
    ])
    doc.add_heading("4. Runtime and IPC", level=1)
    doc.add_paragraph("The Tauri runtime owns the 60 Hz clock, feeds simulator or hardware events into the engine, emits telemetry://snapshot to both windows and sends batched track records to the sidecar. Tauri commands provide snapshot pulls, weather queries, geofence administration, emergency override and intercept calculation. The sidecar uses one JSON object per line over standard input and output.")
    doc.add_heading("5. Limitations and Roadmap", level=1)
    add_doc_bullets(doc, [
        "The real RTL-SDR I/Q and DSP front-end remains the next hardware integration task.",
        "The diagnostics spectrum is a visual prototype until live FFT bins are connected.",
        "BDS register codecs are implemented; live Comm-B interrogation scheduling remains pending.",
        "Operational deployment requires formal safety assessment, frequency coordination, licensing and authority approval.",
    ])


def manual_content(doc: Document) -> None:
    doc.add_heading("1. Purpose and Scope", level=1)
    doc.add_paragraph("This guide covers bench operation using the simulator and optional receive hardware. It is not an operational approval, an ATC procedure or an authorisation to transmit on aviation frequencies.")
    doc.add_heading("2. Installation", level=1)
    add_doc_table(doc, ["Requirement", "Bench recommendation"], [
        ["Host", "Current laptop or desktop; 4 CPU cores and 8 GB RAM recommended"],
        ["Receiver", "Three RTL-SDR Blog V4 units for ADS-B, ACARS and VHF receive"],
        ["Transmitter", "Not required for receive-only bench testing; HackRF TX requires authorisation and controlled test conditions"],
        ["Audio", "Standard USB headset with microphone"],
        ["Power", "Powered USB 3.0 hub; do not run three receivers from a passive hub"],
    ], [1.5, 5.0])
    doc.add_heading("3. Start the Software", level=1)
    add_doc_table(doc, ["Mode", "Command"], [
        ["Full desktop", "scripts/desktop.sh"],
        ["Browser preview", "npm run dev, then open http://localhost:1420/"],
        ["Operations display", "Open http://localhost:1420/hud.html or use the Operations Display button"],
        ["Headless", "cd src-tauri && cargo run --no-default-features --bin aeropulse-headless"],
        ["Verification", "./scripts/verify.sh"],
    ], [1.8, 4.7])
    doc.add_heading("4. Display 1 — Radar", level=1)
    doc.add_paragraph("Select an aircraft by clicking its symbol. Use Measure Distance, click a start point and move the pointer; double-click clears the measurement. Use Operations Display to open the second window. The radar uses kilometre range rings, metre altitude and kilometre-per-hour speed presentation.")
    doc.add_heading("5. Display 2 — Operations Workspace", level=1)
    doc.add_paragraph("On a small screen the navigation bar switches between complete sections: Flight Strips, Aircraft, Weather, Alerts, Audio and System. Only the selected section is displayed in compact mode, preventing one card from covering another. Each section has its own scroll area where required.")
    doc.add_heading("6. Aircraft Details and Audio", level=1)
    doc.add_paragraph("Select a flight strip or radar contact to open its individual details. The detail view shows identification, position, altitude, speed, heading, vertical rate, accuracy and tracking state. The Transmit to aircraft action tunes the audio panel to the assigned prototype channel; it does not bypass the transmission safety interlock.")
    doc.add_heading("7. Safety Rules", level=1)
    add_doc_bullets(doc, [
        "Treat all simulated values as test data until independently verified against a known source.",
        "Receive-only testing is the default and requires no transmitter.",
        "Never connect a transmitter to an antenna or radiate on an aviation frequency without the required Nigerian authorisation and a controlled test plan.",
        "The software uses hold-to-talk, frequency validation, emergency inhibition and a 30-second transmission timeout; these are safeguards, not regulatory approval.",
    ])
    doc.add_heading("8. Troubleshooting", level=1)
    add_doc_table(doc, ["Symptom", "Action"], [
        ["Upper-air table blank", "Select Weather in the compact navigation. If no reports have arrived, the panel shows aircraft estimates or an explicit waiting state."],
        ["Flight strip missing", "Select Flight Strips and scroll inside the strip list; the list is independent of the page."],
        ["Audio output missing", "Allow microphone permission, reconnect the USB headset and choose Primary Output again."],
        ["Desktop shell will not build", "Install webkit2gtk-4.1 and libsoup-3.0 development packages or source scripts/webkit-env.sh."],
    ], [1.8, 4.7])


def technical_content(doc: Document) -> None:
    doc.add_heading("1. Hardware and Wire Formats", level=1)
    doc.add_paragraph("mode_s_decoder.rs accepts 7-byte and 14-byte Mode S frames, validates CRC-24 address/even parity, attempts single-bit repair and decodes DF17 identity, airborne position, velocity and status messages. CPR pairs are resolved inside a 10-second epoch. acars_decoder.rs parses SOH/MODE/TAIL/ACK/LABEL/BLOCK/STX/TEXT/ETX/BCS and classifies METAR, SPECI and D-ATIS reports.")
    doc.add_heading("2. Kinematics", level=1)
    add_doc_table(doc, ["Component", "Contract"], [
        ["EKF", "State [x,y,z,vx,vy,vz], ENU metres and metres per second; Joseph covariance update; acquisition warm-up and Mahalanobis gate"],
        ["Geodesy", "SiteOrigin converts WGS-84 latitude/longitude to local ENU and back"],
        ["R*-tree", "AABB index with overlap-based subtree selection, margin split and forced reinsertion"],
        ["STCA", "5 NM / 1,000 ft thresholds, 120-second lookahead, 5-second refinement"],
    ], [1.5, 5.0])
    doc.add_heading("3. Weather and Defence", level=1)
    doc.add_paragraph("Weather fusion retains AWOS, ACARS and BDS-derived nodes, applies source priority and performs vector-based inverse-distance interpolation with an ISA fallback. Defence modules map 7500, 7600 and 7700; detect identity loss and silence; monitor polygon/floor/ceiling geofences; and solve constant-velocity intercept geometry.")
    doc.add_heading("4. Engine Snapshot", level=1)
    doc.add_paragraph("EngineSnapshot contains tracks, flight data blocks, STCA alerts, geofence breaches, anomaly notes, weather and EngineStatus. The Rust-to-TypeScript field names remain snake_case. The frontend converts wire feet/knots to metre and kilometre-per-hour presentation through src/lib/units.ts.")
    doc.add_heading("5. Frontend Components", level=1)
    add_doc_table(doc, ["Component", "Responsibility"], [
        ["RadarCanvas", "60 fps radar drawing, aircraft selection, range/bearing measurement and conflict lines"],
        ["FlightStripTable", "Arrival/departure strips and aircraft selection"],
        ["AircraftDetailPanel", "Individual aircraft data and audio-channel action"],
        ["WeatherGridHUD", "Surface and upper-air conditions with fallback state"],
        ["AudioRoutingPanel", "Two output devices, frequency selection and guarded PTT"],
        ["SystemMetrics", "Receiver status, spectrum visualisation and performance gauges"],
    ], [1.8, 4.7])
    doc.add_heading("6. Verification Inventory", level=1)
    add_doc_table(doc, ["Suite", "Expected result"], [
        ["cargo test --no-default-features", "90 Rust tests pass"],
        ["python-sidecar pytest", "14 Python tests pass"],
        ["npm run typecheck", "TypeScript compiler exits cleanly"],
        ["npm run build", "Vite emits index.html and hud.html"],
    ], [2.8, 3.7])


def build_pitch() -> None:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    def add_slide(title: str, section: str, accent=ACCENT):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = ppt_colour(PITCH_BLACK)
        # restrained header rule and footer
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(13.333), PptInches(0.08))
        line.fill.solid(); line.fill.fore_color.rgb = ppt_colour(accent); line.line.fill.background()
        footer = slide.shapes.add_textbox(PptInches(0.55), PptInches(7.08), PptInches(12.2), PptInches(0.22))
        footer.text_frame.text = f"AeroPulse-NG  |  NDAIE 2026  |  {section}"
        fp = footer.text_frame.paragraphs[0]; fp.font.size = PptPt(8); fp.font.name = "Aptos"; fp.font.color.rgb = ppt_colour(SILVER_DIM)
        return slide

    def add_text(slide, text, x, y, w, h, size=18, colour=WHITE, bold=False, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        tf = shape.text_frame; tf.word_wrap = True; tf.margin_left = PptInches(0.03); tf.margin_right = PptInches(0.03)
        p = tf.paragraphs[0]; p.text = text; p.alignment = align
        p.font.size = PptPt(size); p.font.bold = bold; p.font.name = "Aptos"; p.font.color.rgb = ppt_colour(colour)
        return shape

    # 1 Title
    slide = add_slide("", "Title", CIVIL)
    logo = ASSETS / "logo.png"
    if logo.exists(): slide.shapes.add_picture(str(logo), PptInches(0.8), PptInches(1.2), width=PptInches(1.35))
    add_text(slide, "AeroPulse-NG", 2.35, 1.35, 9.8, 0.7, 38, WHITE, True)
    add_text(slide, "Offline air surveillance for Nigerian operations", 2.38, 2.15, 9.5, 0.45, 20, CIVIL, True)
    add_text(slide, "National Drone & Aerospace Innovation Expo 2026\nNational Innovation Pitch Challenge", 2.38, 3.1, 9, 0.9, 16, SILVER)
    add_text(slide, "Primary category: Flight Planning & Air Traffic Management", 2.38, 4.35, 9, 0.4, 14, SILVER_DIM)
    add_text(slide, "Air-gapped  •  Offline-first  •  SI metric presentation", 2.38, 5.3, 9, 0.4, 13, SILVER_DIM)

    # 2 Problem
    slide = add_slide("", "The Problem", CRITICAL)
    add_text(slide, "A surveillance gap remains below and between radar sites", 0.7, 0.65, 12, 0.65, 28, WHITE, True)
    points = [
        ("Cost", "Conventional primary and secondary radar installations require major capital and specialist maintenance."),
        ("Coverage", "Low-altitude corridors, secondary aerodromes and terrain-shadowed sectors remain difficult to monitor."),
        ("Resilience", "Loss of connectivity or a radar head removes the traffic picture during periods of elevated risk."),
        ("Cooperation", "Transponder-silent or unidentified aircraft are difficult to track with dependent surveillance alone."),
    ]
    y = 1.8
    for head, body in points:
        add_text(slide, head, 1.0, y, 1.5, 0.35, 17, CRITICAL, True)
        add_text(slide, body, 2.5, y, 9.7, 0.65, 15, SILVER)
        y += 1.1

    # 3 Solution
    slide = add_slide("", "The Solution", CIVIL)
    add_text(slide, "A local surveillance workstation from affordable receiver hardware", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    if (ASSETS / "radar-scope.png").exists(): slide.shapes.add_picture(str(ASSETS / "radar-scope.png"), PptInches(0.7), PptInches(1.55), width=PptInches(7.2))
    add_text(slide, "AeroPulse-NG combines:", 8.2, 1.75, 4.4, 0.35, 18, CIVIL, True)
    bullets = ["1090 MHz ADS-B and Mode S decoding", "Filtered tracks with dead reckoning", "Conflict alerting within a 120-second lookahead", "Offline weather fusion for local conditions", "A dual-display operator workspace", "No cloud service required"]
    y = 2.25
    for item in bullets:
        add_text(slide, "• " + item, 8.25, y, 4.2, 0.45, 13, SILVER); y += 0.62

    # 4 Architecture
    slide = add_slide("", "How It Works", CIVIL)
    add_text(slide, "One processing chain from signal to decision support", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    if (ASSETS / "architecture.png").exists(): slide.shapes.add_picture(str(ASSETS / "architecture.png"), PptInches(0.55), PptInches(1.45), width=PptInches(12.2))

    # 5 Engineering
    slide = add_slide("", "Engineering", MILV)
    add_text(slide, "Engineering depth is measurable, not decorative", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    blocks = [("Signal", ["DF17 CRC validation and repair", "CPR even/odd position resolution", "ACARS message integrity"]), ("Track", ["6-state Extended Kalman Filter", "R*-tree spatial indexing", "Dead-reckoning during dropouts"]), ("Safety", ["9.26 km / 305 m separation view", "120-second lookahead", "Emergency and silence alerts"]), ("Weather", ["AWOS + ACARS + BDS inputs", "Vector wind interpolation", "Standard-atmosphere fallback"])]
    x = 0.7
    for heading, items in blocks:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(1.65), PptInches(2.95), PptInches(4.55))
        shape.fill.solid(); shape.fill.fore_color.rgb = ppt_colour(OBSIDIAN); shape.line.color.rgb = ppt_colour(SILVER_DIM)
        add_text(slide, heading, x+0.18, 1.9, 2.55, 0.35, 16, CIVIL, True)
        yy = 2.55
        for item in items:
            add_text(slide, "• " + item, x+0.18, yy, 2.55, 0.65, 12, SILVER); yy += 0.82
        x += 3.15
    add_text(slide, "Prototype verification: 90 Rust tests  •  14 Python tests  •  strict TypeScript build", 0.75, 6.45, 11.7, 0.35, 13, SILVER_DIM)

    # 6 Demo
    slide = add_slide("", "Demonstration", WARNING)
    add_text(slide, "A scenario the judges can observe", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    if (ASSETS / "timeline.png").exists(): slide.shapes.add_picture(str(ASSETS / "timeline.png"), PptInches(0.7), PptInches(1.55), width=PptInches(12))
    add_text(slide, "The simulator exercises the same event path used by the future receiver integration.", 0.9, 6.05, 11.5, 0.4, 14, SILVER)

    # 7 Weather
    slide = add_slide("", "Weather", WARNING)
    add_text(slide, "Three local sources, one weather picture", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    if (ASSETS / "fusion.png").exists(): slide.shapes.add_picture(str(ASSETS / "fusion.png"), PptInches(1.0), PptInches(1.55), width=PptInches(11.3))
    add_text(slide, "Surface observations take priority; sparse upper-air observations fall back to the standard atmosphere with an explicit confidence value.", 1.0, 6.0, 11.3, 0.45, 14, SILVER)

    # 8 Standards
    slide = add_slide("", "Standards", CIVIL)
    add_text(slide, "Standards and presentation policy", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    rows = [("Units", "NCAA / ICAO Annex 5", "km · m · km/h · hPa"), ("Surveillance", "DO-260B", "Mode S / ADS-B / CPR"), ("Separation", "ICAO Doc 4444", "9.26 km · 305 m · 120 s"), ("Display", "HF-STD-010A", "Readable, consistent data blocks"), ("Symbology", "MIL-STD-2525D", "Consistent civil/military frames"), ("Integrity", "NDAIE Rulebook", "AI assistance disclosed")]
    y = 1.8
    for a, b, c in rows:
        add_text(slide, a, 1.0, y, 1.5, 0.35, 15, CIVIL, True)
        add_text(slide, b, 2.7, y, 4.3, 0.35, 14, SILVER)
        add_text(slide, c, 7.7, y, 4.3, 0.35, 14, SILVER_DIM)
        y += 0.72

    # 9 Impact
    slide = add_slide("", "Impact", CIVIL)
    add_text(slide, "A practical path to local capability", 0.7, 0.65, 12, 0.65, 27, WHITE, True)
    cards = [(">99%", "potential capital reduction per station", CIVIL), ("24/7", "local decision support without cloud dependence", MILV), ("SI", "metric-first display for Nigerian operators", CIVIL2), ("Local", "maintainable engineering and training platform", WARNING)]
    x = 0.8
    for big, small, col in cards:
        add_text(slide, big, x, 1.8, 2.8, 0.7, 30, col, True, PP_ALIGN.CENTER)
        add_text(slide, small, x+0.2, 2.7, 2.4, 1.0, 14, SILVER, False, PP_ALIGN.CENTER)
        x += 3.1
    add_text(slide, "The prototype is designed for progressive validation: simulator → bench receivers → controlled field trial.", 1.0, 5.3, 11.3, 0.5, 18, SILVER)

    # 10 Close
    slide = add_slide("", "Next Steps", CIVIL)
    add_text(slide, "Ready for the next validation stage", 0.7, 0.9, 12, 0.65, 30, WHITE, True)
    add_text(slide, "1", 1.1, 2.0, 0.5, 0.5, 24, CIVIL, True)
    add_text(slide, "Purchase the bench receiver kit", 1.8, 2.0, 5, 0.45, 18, SILVER, True)
    add_text(slide, "2", 1.1, 3.0, 0.5, 0.5, 24, CIVIL, True)
    add_text(slide, "Validate ADS-B, ACARS and receive-only VHF paths", 1.8, 3.0, 7, 0.45, 18, SILVER, True)
    add_text(slide, "3", 1.1, 4.0, 0.5, 0.5, 24, CIVIL, True)
    add_text(slide, "Complete controlled hardware and safety assessment before any transmission", 1.8, 4.0, 9.5, 0.7, 18, SILVER, True)
    add_text(slide, "AeroPulse-NG  |  Nigerian airspace, locally supported", 1.1, 5.7, 10, 0.45, 17, CIVIL)

    prs.save(OUT / "AeroPulse-NG_Pitch_Deck.pptx")


def main() -> None:
    build_pitch()
    build_docx("AeroPulse-NG_Executive_Summary.docx", "AeroPulse-NG — Executive Project Summary", "NDAIE 2026  |  Flight Planning & Air Traffic Management", executive_content)
    build_docx("AeroPulse-NG_Architecture.docx", "AeroPulse-NG — System Architecture", "Design, data flow and implementation boundaries", architecture_content)
    build_docx("AeroPulse-NG_User_Manual.docx", "AeroPulse-NG — User Manual", "Bench setup, operation and safety guidance", manual_content)
    build_docx("AeroPulse-NG_Technical_Reference.docx", "AeroPulse-NG — Technical Reference", "Wire contracts, algorithms and verification", technical_content)
    print(f"Office documents written to {OUT}")


if __name__ == "__main__":
    main()
